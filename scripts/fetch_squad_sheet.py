#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Zaciąga PIERWSZY SKŁAD Rakowa z arkusza Google (który klub aktualizuje) i zapisuje
public/squad.json. Uruchamiane w pipeline PRZED discover_statsbomb.py, żeby skład
w aplikacji nadążał za rzeczywistością bez ręcznej edycji.

Źródło prawdy: arkusz z kolumnami [Drużyna, Imię, Nazwisko, Pozycja, ...]. Pierwszy
skład = wiersze, gdzie „Drużyna" == „Raków Częstochowa" (dokładnie; „…II", „…U19",
„…Youth" są pomijane — to jest „filtr pierwszego składu").

Zasady bezpieczeństwa (żeby zła pobrana wersja nie wywaliła składu):
  • pobranie się nie uda LUB pierwszy skład < MIN_SQUAD → NIE dotykamy squad.json;
  • dopasowanie do istniejącego squad.json po nazwisku (rozmyte) ZACHOWUJE zapis
    nazwy i id znany StatsBombowi (inaczej „Adriano Luís…" zerwałoby dopasowanie
    RC dla „Adriano"). Automat zarządza SKŁADEM (dodaj/usuń), nie przepisuje pozycji
    utrwalonych zawodników — rozbieżność pozycji tylko loguje do przeglądu analityka.

Użycie:
  python scripts/fetch_squad_sheet.py                # pobiera z SHEET_URL
  python scripts/fetch_squad_sheet.py plik.xlsx      # tryb testowy z pliku lokalnego
  SQUAD_SHEET_URL=... python scripts/fetch_squad_sheet.py

Wymaga: openpyxl. Sieć: requests (jeśli brak — urllib z stdlib).
"""
import sys, os, io, json, re, unicodedata
from pathlib import Path

# ID arkusza wystawionego przez Raków. Arkusz musi być udostępniony „każdy z linkiem
# może wyświetlać", żeby pipeline mógł go pobrać bez logowania. Można nadpisać zmienną
# środowiskową SQUAD_SHEET_URL (pełny URL eksportu) albo SQUAD_SHEET_ID.
SHEET_ID = os.environ.get("SQUAD_SHEET_ID") or "13DgdcL9OM0L1_47YBIEDqtXmSDf7R5mn"
SHEET_URL = os.environ.get("SQUAD_SHEET_URL") \
    or f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

FIRST_TEAM = os.environ.get("SQUAD_FIRST_TEAM") or "Raków Częstochowa"  # dokładna nazwa drużyny
MIN_SQUAD = 15           # mniej niż tyle = coś poszło nie tak, nie nadpisuj
OUT = Path(__file__).resolve().parent.parent / "public" / "squad.json"

# Mapowanie pozycji z arkusza na taksonomię modelu (pos + linia).
POSMAP = {
    "goalkeeper": ("GK", "Bramka"),
    "centerback": ("CB", "Obrona"), "centreback": ("CB", "Obrona"),
    "leftback": ("WM", "Pomoc"), "rightback": ("WM", "Pomoc"), "wingback": ("WM", "Pomoc"),
    "leftwingback": ("WM", "Pomoc"), "rightwingback": ("WM", "Pomoc"),
    "defensivemidfield": ("DM", "Pomoc"), "centralmidfield": ("DM", "Pomoc"),
    "attackingmidfield": ("AM", "Pomoc"),
    "leftmidfield": ("WM", "Pomoc"), "rightmidfield": ("WM", "Pomoc"),
    "leftwinger": ("WM", "Pomoc"), "rightwinger": ("WM", "Pomoc"), "winger": ("WM", "Pomoc"),
    "centerforward": ("ST", "Atak"), "centreforward": ("ST", "Atak"),
    "secondstriker": ("ST", "Atak"), "forward": ("ST", "Atak"), "striker": ("ST", "Atak"),
}
LINE_ORDER = {"Bramka": 0, "Obrona": 1, "Pomoc": 2, "Atak": 3}

def fold(s):
    """Nazwa → porównywalna: bez diakrytyków (ł→l), małe litery, myślnik→spacja."""
    s = unicodedata.normalize("NFD", str(s or "")).replace("ł", "l").replace("Ł", "l")
    s = s.encode("ascii", "ignore").decode().lower().replace("-", " ")
    return " ".join(s.split())
def toks(s):
    return set(fold(s).split())
def slug(name):
    s = str(name or "").strip().lower().replace(" ", "-")
    return "rk-" + re.sub(r"[^0-9a-zżźćńśłóęąóa-z\-]", "", s)

def load_bytes(src):
    """Zwraca bajty xlsx: z pliku lokalnego albo z URL (requests → fallback urllib)."""
    if os.path.exists(src):
        return Path(src).read_bytes()
    try:
        import requests
        r = requests.get(src, timeout=45, allow_redirects=True)
        r.raise_for_status()
        return r.content
    except ImportError:
        import urllib.request
        with urllib.request.urlopen(src, timeout=45) as resp:
            return resp.read()

def parse_first_team(xlsx_bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    ws = wb["Skład Rakowa"] if "Skład Rakowa" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    hdr = [str(c or "").strip().lower() for c in rows[0]]
    def col(*names):
        for n in names:
            if n in hdr:
                return hdr.index(n)
        return None
    ci = {"team": col("drużyna", "druzyna", "team"), "im": col("imię", "imie", "first name"),
          "nz": col("nazwisko", "last name"), "poz": col("pozycja", "position")}
    if ci["team"] is None or ci["nz"] is None:
        raise ValueError(f"Nie rozpoznano kolumn w arkuszu (nagłówki: {hdr[:12]})")
    out = []
    for r in rows[1:]:
        team = str(r[ci["team"]] or "").strip()
        if fold(team) != fold(FIRST_TEAM):
            continue
        nm = " ".join(x for x in [str(r[ci["im"]] or "").strip() if ci["im"] is not None else "",
                                   str(r[ci["nz"]] or "").strip()] if x).strip()
        if not nm:
            continue
        poz = str(r[ci["poz"]] or "").strip().lower().replace(" ", "") if ci["poz"] is not None else ""
        out.append({"name": nm, "poz_raw": poz})
    return out

def main():
    src = sys.argv[1] if len(sys.argv) > 1 else SHEET_URL
    # 1) Pobierz + sparsuj — każdy błąd = zostaw squad.json bez zmian.
    try:
        players = parse_first_team(load_bytes(src))
    except Exception as e:
        print(f"[squad-sheet] Nie udało się pobrać/odczytać arkusza ({e}). "
              f"Zostawiam istniejący squad.json.", file=sys.stderr)
        return 0
    if len(players) < MIN_SQUAD:
        print(f"[squad-sheet] Pierwszy skład z arkusza ma tylko {len(players)} "
              f"zawodników (< {MIN_SQUAD}) — podejrzane. Nie nadpisuję squad.json.", file=sys.stderr)
        return 0

    # 2) Istniejący squad.json — zachowujemy nazwy/id znane StatsBombowi.
    try:
        existing = json.loads(OUT.read_text(encoding="utf-8"))
    except Exception:
        existing = []
    ex_tok = [(x, toks(x.get("name", ""))) for x in existing]

    def find_existing(nm):
        t = toks(nm)
        best = None
        for x, xt in ex_tok:
            if t <= xt or xt <= t or len(t & xt) >= 2:
                # preferuj najwięcej wspólnych tokenów
                score = len(t & xt)
                if best is None or score > best[1]:
                    best = (x, score)
        return best[0] if best else None

    squad, added, removed, pos_warn = [], [], [], []
    kept_names = set()
    for pl in players:
        ex = find_existing(pl["name"])
        pm = POSMAP.get(pl["poz_raw"])
        if ex:
            kept_names.add(ex.get("name"))
            entry = {"id": ex.get("id") or slug(ex["name"]), "name": ex["name"],
                     "pos": ex.get("pos"), "line": ex.get("line")}
            # Rozbieżność pozycji arkusz vs nasze — tylko log do przeglądu (nie zmieniamy).
            if pm and ex.get("pos") and pm[0] != ex["pos"]:
                pos_warn.append(f"{ex['name']}: arkusz={pl['poz_raw']}→{pm[0]}, u nas {ex['pos']}")
        else:
            if not pm:
                print(f"[squad-sheet] Nowy zawodnik '{pl['name']}' z pozycją '{pl['poz_raw']}' "
                      f"bez mapowania — pomijam (uzupełnij POSMAP).", file=sys.stderr)
                continue
            entry = {"id": slug(pl["name"]), "name": pl["name"], "pos": pm[0], "line": pm[1]}
            added.append(pl["name"])
        squad.append(entry)

    for x in existing:
        if x.get("name") not in kept_names:
            removed.append(x.get("name"))

    squad.sort(key=lambda e: (LINE_ORDER.get(e["line"], 9), e["name"]))
    OUT.write_text(json.dumps(squad, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"[squad-sheet] Pierwszy skład z arkusza: {len(players)} → zapisano {len(squad)} do squad.json.")
    if added:   print(f"[squad-sheet] Dodani: {', '.join(added)}")
    if removed: print(f"[squad-sheet] Usunięci (poza pierwszym składem): {', '.join(removed)}")
    if pos_warn:
        print("[squad-sheet] Rozbieżności pozycji do przeglądu:")
        for w in pos_warn: print(f"    - {w}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
